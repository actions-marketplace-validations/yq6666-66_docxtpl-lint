"""Build realistic templates for the self-evaluation (docs/eval/self-eval.md).

python scripts/realworld.py OUTDIR  ->  invoice.docx, 合同 模板.docx, big.docx, invoice_v2.docx, broken.docx, data*.json
"""

import json
import sys
from pathlib import Path

from docx import Document

out = Path(sys.argv[1] if len(sys.argv) > 1 else "rw")
out.mkdir(parents=True, exist_ok=True)


def invoice(path, v2=False):
    d = Document()
    d.sections[0].header.paragraphs[0].text = "{{ company.name }} · {{ company.vat_id }}"
    d.sections[0].footer.paragraphs[0].text = "Invoice {{ invoice.no }} — {{ invoice.date }}"
    d.add_paragraph("Bill to: {{ customer.name }}")
    d.add_paragraph("{{ customer.address.street }}, {{ customer.address.city }}")
    t = d.add_table(rows=4, cols=4)
    for j, h in enumerate(["Item", "Qty", "Price", "Total"]):
        t.cell(0, j).text = h
    t.cell(1, 0).text = "{%tr for line in lines %}"
    t.cell(2, 0).text = "{{ line.name }}"
    t.cell(2, 1).text = "{{ line.qty }}"
    t.cell(2, 2).text = "{{ '%.2f' % line.price }}"
    t.cell(2, 3).text = "{{ (line.qty * line.price)|round(2) }}"
    t.cell(3, 0).text = "{%tr endfor %}"
    d.add_paragraph("{%p if discount is defined and discount > 0 %}")
    d.add_paragraph("Discount: {{ discount }} %")
    d.add_paragraph("{%p endif %}")
    d.add_paragraph("Total: {{ total|round(2) }} {{ currency|default('EUR') }}")
    if v2:
        d.add_paragraph("PO number: {{ invoice.po_number }}")
        d.add_paragraph("{% for n in notes %}{{ n }}{% endfor %}")
    d.save(path)


def contract_zh(path):
    d = Document()
    d.sections[0].header.paragraphs[0].text = "合同编号：{{ 合同.编号 if false else contract.no }}"
    d.add_paragraph("甲方：{{ party_a.name }}（统一社会信用代码：{{ party_a.uscc }}）")
    d.add_paragraph("乙方：{{ party_b.name }}")
    d.add_paragraph("{%p for clause in clauses %}")
    d.add_paragraph("第{{ loop.index }}条 {{ clause.title }}：{{ clause.body }}")
    d.add_paragraph("{%p endfor %}")
    outer = d.add_table(rows=2, cols=2)
    outer.cell(0, 0).text = "付款计划"
    inner = outer.cell(0, 1).add_table(rows=3, cols=2)
    inner.cell(0, 0).text = "{%tr for p in payments %}"
    inner.cell(1, 0).text = "{{ p.date }}"
    inner.cell(1, 1).text = "{% cellbg p.color %}{{ p.amount }} 元"
    inner.cell(2, 0).text = "{%tr endfor %}"
    outer.cell(1, 0).text = "签字"
    outer.cell(1, 1).text = "{{ signer|default('（待定）') }}"
    d.add_paragraph("“{{ party_a.name }}”与“{{ party_b.name }}”签署。")  # smart quotes outside tags
    d.save(path)


def broken(path):
    d = Document()
    d.sections[0].footer.paragraphs[0].text = "{% if show_page %}Page"
    d.add_paragraph("Dear {{ customer.name }},")
    t = d.add_table(rows=3, cols=2)
    t.cell(0, 0).text = "{% tr for i in items %}"
    t.cell(1, 0).text = "{{ i.name }}"
    t.cell(1, 1).text = "{{ i.price }"
    t.cell(2, 0).text = "{%tr endfro %}"
    d.add_paragraph("Total: {{ total }} {%p if vat %}")
    d.add_paragraph("{%p endif %}")
    d.save(path)


def big(path, n=3000, rows=300):
    d = Document()
    for i in range(n):
        d.add_paragraph(f"Clause {i}: {{{{ c{i % 50}.text }}}} {{% if c{i % 50}.on %}}on{{% endif %}}")
    t = d.add_table(rows=rows + 2, cols=3)
    t.cell(0, 0).text = "{%tr for r in rows %}"
    for j in range(3):
        t.cell(1, j).text = f"{{{{ r.f{j} }}}}"
    t.cell(rows + 1, 0).text = "{%tr endfor %}"
    d.save(path)


invoice(out / "invoice.docx")
invoice(out / "invoice_v2.docx", v2=True)
contract_zh(out / "合同 模板.docx")
broken(out / "broken.docx")
big(out / "big.docx")
data = {
    "company": {"name": "ACME", "vat_id": "DE1"},
    "invoice": {"no": "2026-001", "date": "2026-10-16"},
    "customer": {"name": "Zoë", "address": {"street": "Main 1", "city": "Köln"}},
    "lines": [{"name": "Widget", "qty": 2, "price": 9.5}],
    "total": 19.0,
}
(out / "data.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
# the same file as Windows Notepad saves it: UTF-8 with BOM and CRLF
(out / "data_bom_crlf.json").write_bytes(
    b"\xef\xbb\xbf" + json.dumps(data, ensure_ascii=False, indent=1).replace("\n", "\r\n").encode()
)
print("built", sorted(p.name for p in out.iterdir()))
