"""v1.2: ideas adopted from the OSS survey (docs/research/oss-survey.md), verified by real docxtpl renders."""

import json
import zipfile

import pytest
from docx import Document
from docxtpl import DocxTemplate

from docxtpl_lint.cli import main
from docxtpl_lint.contract import infer_tree, sample, to_pydantic
from docxtpl_lint.core import lint_path


def _doc(tmp_path, *paras, name="t.docx"):
    d = Document()
    for p in paras:
        d.add_paragraph(p)
    path = tmp_path / name
    d.save(path)
    return path


def _render(path, ctx):
    t = DocxTemplate(str(path))
    t.render(ctx)
    out = path.with_name("out_" + path.name)
    t.save(out)
    return [p.text for p in Document(out).paragraphs]


# ---- F1/F2 paragraph-level suppression + unused suppression report (ruff noqa / RUF100)

PCT = "{%p if n % 2 == 0 %}"


def test_disable_line(tmp_path):
    p = _doc(tmp_path, PCT + " {# docxtpl-lint: disable-line=DTL201 #}", "x", "{%p endif %}")
    assert "DTL201" not in [i.rule for i in lint_path(p).issues]
    q = _doc(tmp_path, PCT, "x", "{%p endif %}", name="q.docx")
    assert "DTL201" in [i.rule for i in lint_path(q).issues]


def test_disable_next_with_p_comment_is_removed_by_docxtpl(tmp_path):
    p = _doc(tmp_path, "A", "{#p docxtpl-lint: disable-next=DTL209 #}", "lost {%p if x %}", "B", "{%p endif %}")
    rules = [i.rule for i in lint_path(p).issues]
    assert "DTL209" not in rules and "DTL005" not in rules
    # docxtpl drops the whole {#p ... #} paragraph: the directive leaves no blank line
    q = _doc(tmp_path, "A", "{#p docxtpl-lint: disable-next=DTL201 #}", "{{ v }}", name="q.docx")
    assert _render(q, {"v": "1"}) == ["A", "1"]


def test_disable_next_only_covers_next_paragraph(tmp_path):
    p = _doc(tmp_path, "{#p docxtpl-lint: disable-next=DTL209 #}", "ok", "lost {%p if x %}", "{%p endif %}")
    rules = [i.rule for i in lint_path(p).issues]
    assert "DTL209" in rules and "DTL005" in rules


def test_unused_and_unknown_suppression(tmp_path, capsys):
    p = _doc(tmp_path, "{{ a }} {# docxtpl-lint: disable-line=DTL201,DTL999 #}")
    iss = lint_path(p).issues
    assert [i.rule for i in iss] == ["DTL005"]
    assert "DTL201, DTL999 (unknown rule)" in iss[0].message
    assert "DTL999（未知规则）" in iss[0].message_for("zh")


def test_context_rule_suppression_not_reported_unused(tmp_path):
    p = _doc(tmp_path, "{{ a }} {# docxtpl-lint: disable-line=DTL301 #}")
    assert lint_path(p).issues == []
    (tmp_path / "d.json").write_text("{}")
    assert main([str(p), "--data", str(tmp_path / "d.json"), "--no-config"]) == 0


def test_prefix_rule_in_suppression(tmp_path):
    p = _doc(tmp_path, PCT + " {# docxtpl-lint: disable-line=DTL2 #}", "x", "{%p endif %}")
    assert [i.rule for i in lint_path(p).issues] == []


def test_overlapping_rules_in_one_suppression(tmp_path):
    p = _doc(tmp_path, PCT + " {# docxtpl-lint: disable-line=DTL2,DTL201 #}", "x", "{%p endif %}")
    assert [i.rule for i in lint_path(p).issues] == []


# ---- F3 --statistics


def test_statistics(tmp_path, capsys):
    p = _doc(tmp_path, "{% if a %}", "{% for x in y %}", "{%tr if b %}")
    assert main([str(p), "--statistics", "--no-config"]) == 1
    out = capsys.readouterr().out.splitlines()
    assert out[0].split()[:2] == ["3", "DTL101"] and any("DTL203" in line for line in out)


# ---- F4 baseline


def test_baseline_roundtrip_and_paragraph_shift(tmp_path, capsys):
    p = _doc(tmp_path, "intro", "{% if a %}")
    bl = tmp_path / "bl.json"
    assert main([str(p), "--write-baseline", str(bl), "--no-config"]) == 0
    data = json.loads(bl.read_text())
    assert data["version"] == 1 and list(data["files"]) == ["t.docx"]
    assert main([str(p), "--baseline", str(bl), "--no-config"]) == 0
    # paragraphs inserted above: the same finding is still baselined
    _doc(tmp_path, "new", "more", "intro", "{% if a %}")
    assert main([str(p), "--baseline", str(bl), "--no-config"]) == 0
    # a new error fails; the old one stays hidden
    _doc(tmp_path, "intro", "{% if a %}", "{{ b }")
    capsys.readouterr()
    assert main([str(p), "--baseline", str(bl), "--no-config", "--format", "json"]) == 1
    rules = [i["rule"] for f in json.loads(capsys.readouterr().out)["files"] for i in f["issues"]]
    assert rules == ["DTL001"]


def test_baseline_stale_reported(tmp_path, capsys):
    p = _doc(tmp_path, "{% if a %}")
    bl = tmp_path / "bl.json"
    main([str(p), "--write-baseline", str(bl), "--no-config"])
    _doc(tmp_path, "{% if a %}{% endif %}")
    capsys.readouterr()
    assert main([str(p), "--baseline", str(bl), "--no-config"]) == 0
    assert "1 stale baseline entry" in capsys.readouterr().err


def test_baseline_bad_file(tmp_path, capsys):
    p = _doc(tmp_path, "x")
    (tmp_path / "bl.json").write_text("[]")
    assert main([str(p), "--baseline", str(tmp_path / "bl.json"), "--no-config"]) == 2


def test_baseline_from_pyproject(tmp_path, monkeypatch):
    p = _doc(tmp_path, "{% if a %}")
    main([str(p), "--write-baseline", str(tmp_path / "bl.json"), "--no-config"])
    (tmp_path / "pyproject.toml").write_text('[tool.docxtpl-lint]\nbaseline = "bl.json"\n')
    monkeypatch.chdir(tmp_path)
    assert main([str(p)]) == 0


# ---- F6 annotated copy


def test_annotate_writes_comments_and_keeps_original(tmp_path):
    d = Document()
    d.add_paragraph("ok")
    t = d.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "{{ r.price }} {%tr endfor %}"
    d.add_paragraph("{% if a %}")
    src = tmp_path / "inv.docx"
    d.save(src)
    before = src.read_bytes()
    assert main([str(src), "--annotate", str(tmp_path / "ann"), "--no-config", "--lang", "zh"]) == 1
    assert src.read_bytes() == before
    out = tmp_path / "ann" / "inv.annotated.docx"
    with zipfile.ZipFile(out) as zf:
        comments = zf.read("word/comments.xml").decode()
        assert "DTL209" in comments and "DTL101" in comments and "表格行" in comments
        assert "comments+xml" in zf.read("[Content_Types].xml").decode()
        assert "comments.xml" in zf.read("word/_rels/document.xml.rels").decode()
    # still a valid package for python-docx, same text
    assert [p.text for p in Document(out).paragraphs] == [p.text for p in Document(src).paragraphs]


def test_annotate_skips_clean_files(tmp_path):
    p = _doc(tmp_path, "{{ a }}")
    assert main([str(p), "--annotate", str(tmp_path / "ann"), "--no-config"]) == 0
    assert not (tmp_path / "ann" / "t.annotated.docx").exists()


# ---- F7 Pydantic models, validated by a real render


def test_pydantic_models_render(tmp_path):
    pytest.importorskip("pydantic")
    d = Document()
    d.add_paragraph("{{ customer.name }} {{ customer['tax-id'] }} {% if note %}{{ note }}{% endif %}")
    t = d.add_table(rows=3, cols=1)
    t.cell(0, 0).text = "{%tr for it in items %}"
    t.cell(1, 0).text = "{{ it.title }} {{ it.price | round(2) }}"
    t.cell(2, 0).text = "{%tr endfor %}"
    p = tmp_path / "inv.docx"
    d.save(p)
    tree = infer_tree([p])
    code = to_pydantic(tree)
    ns: dict = {}
    exec(compile(code, "models.py", "exec"), ns)
    Model = ns["TemplateData"]
    Model.model_rebuild(_types_namespace=ns)
    data = sample(tree)
    m = Model.model_validate(data)
    from docxtpl_lint.contract import to_schema

    req = {k for k, f in Model.model_fields.items() if f.is_required()}
    assert req == set(to_schema(infer_tree([p]))["required"])  # same contract as the JSON Schema
    out = _render(p, m.model_dump(by_alias=True))
    assert "[customer.name]" in out[0]
    from pydantic import ValidationError

    with pytest.raises(ValidationError):
        Model.model_validate({"items": [{"title": "x", "price": "not a number"}], "customer": {}})


def test_infer_pydantic_cli(tmp_path, capsys):
    p = _doc(tmp_path, "{{ class }} {{ a.b }}")
    assert main(["infer", str(p), "--pydantic", "-"]) == 0
    out = capsys.readouterr().out
    assert "class TemplateData(BaseModel)" in out and "alias='class'" in out
    assert main(["infer", str(p), "--pydantic", str(tmp_path / "m.py"), "-o", str(tmp_path / "s.json")]) == 0
    assert (tmp_path / "m.py").read_text().startswith('"""Generated by docxtpl-lint')


# ---- bug fix: DTL204 Chinese location


def test_dtl204_zh_location_localised(tmp_path):
    d = Document()
    t = d.add_table(rows=2, cols=1)
    t.cell(0, 0).text = "{%tr for x in xs %}"
    t.cell(1, 0).text = "{% endfor %}"
    p = tmp_path / "t.docx"
    d.save(p)
    iss = [i for i in lint_path(p).issues if i.rule == "DTL204"]
    assert iss
    zh = iss[0].message_for("zh")
    assert "表1 行1 列1" in zh and "T1 R1 C1" not in zh
    assert "T1 R1 C1" in iss[0].message_for("en")
