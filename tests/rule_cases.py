"""One bad + one good template per docxtpl-specific rule, with what docxtpl 0.20.2
actually does when rendering the bad one ("raises" or "silent": renders but the
output is wrong). Used by test_docxtpl_behaviour.py and test_rules.py."""

from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass

from docx import Document

CTX = {"xs": [1, 2], "v": 4, "x": True}


@dataclass
class Case:
    name: str
    build: Callable
    rules: list[str]  # expected docxtpl-lint rules
    docxtpl: str  # "ok" | "raises" | "silent"


def _table(d, rows, cols, cells):
    t = d.add_table(rows=rows, cols=cols)
    for (r, c), text in cells.items():
        t.cell(r, c).text = text
    return t


CASES = [
    # DTL201 (#644)
    Case(
        "dtl201_bad",
        lambda d: [d.add_paragraph("{%p if v % 2 == 0 %}"), d.add_paragraph("{%p endif %}")],
        ["DTL201"],
        "raises",
    ),
    Case(
        "dtl201_good",
        lambda d: [
            d.add_paragraph("{% set even = v % 2 == 0 %}"),
            d.add_paragraph("{%p if even %}"),
            d.add_paragraph("{%p endif %}"),
        ],
        [],
        "ok",
    ),
    # DTL202
    Case(
        "dtl202_bad",
        lambda d: _table(d, 3, 1, {(0, 0): "{% tr for i in xs %}", (1, 0): "{{ i }}", (2, 0): "{%tr endfor %}"}),
        ["DTL202"],
        "raises",
    ),
    Case(
        "dtl202_bad_p",
        lambda d: [d.add_paragraph("{% p if x %}"), d.add_paragraph("{%p endif %}")],
        ["DTL202"],
        "raises",
    ),
    # DTL203
    Case(
        "dtl203_bad",
        lambda d: [d.add_paragraph("{%tr for i in xs %}"), d.add_paragraph("{%tr endfor %}")],
        ["DTL203", "DTL203"],
        "raises",
    ),
    # DTL204 (#603)
    Case(
        "dtl204_bad",
        lambda d: _table(d, 3, 1, {(0, 0): "{% for i in xs %}", (1, 0): "{{ i }}", (2, 0): "{%tr endfor %}"}),
        ["DTL204"],
        "silent",
    ),
    # DTL205
    Case(
        "dtl205_bad",
        lambda d: [_table(d, 1, 1, {(0, 0): "{%tr for i in xs %}"}), _table(d, 1, 1, {(0, 0): "{%tr endfor %}"})],
        ["DTL205"],
        "silent",
    ),
    # DTL206
    Case(
        "dtl206_bad",
        lambda d: _table(d, 1, 2, {(0, 0): "{%tr for i in xs %}{{ i }}", (0, 1): "{%tr endfor %}"}),
        ["DTL206"],
        "raises",
    ),
    Case(
        "tr_good",
        lambda d: _table(d, 3, 1, {(0, 0): "{%tr for i in xs %}", (1, 0): "{{ i }}", (2, 0): "{%tr endfor %}"}),
        [],
        "ok",
    ),
    # DTL207
    Case(
        "dtl207_bad",
        lambda d: _table(d, 2, 2, {(0, 0): "{%tc for i in xs %}", (0, 1): "{{ i }}", (1, 1): "{%tc endfor %}"}),
        ["DTL207"],
        "silent",
    ),
    Case(
        "tc_good",
        lambda d: _table(d, 1, 3, {(0, 0): "{%tc for i in xs %}", (0, 1): "{{ i }}", (0, 2): "{%tc endfor %}"}),
        [],
        "ok",
    ),
    # Jinja features docxtpl's default Environment does NOT enable
    Case("do_ext", lambda d: d.add_paragraph("{% set l = [] %}{% do l.append(1) %}{{ l }}"), ["DTL103"], "raises"),
    Case("break_ext", lambda d: d.add_paragraph("{% for i in xs %}{% break %}{% endfor %}"), ["DTL103"], "raises"),
    # DTL208: docxtpl cell tags (found by the docxtpl test-template benchmark)
    Case("dtl208_bad", lambda d: d.add_paragraph("{% cellbg 'FF0000' %}red"), ["DTL208"], "raises"),
    Case(
        "cell_tags_ok",
        lambda d: _table(d, 1, 2, {(0, 0): "{% cellbg 'FF0000' %}a", (0, 1): "{% colspan 1 %}b"}),
        [],
        "ok",
    ),
    Case(
        "smart_quotes_ok",
        lambda d: d.add_paragraph("{{ v | default(\u2018x\u2019) }}{% if \u201cx\u201d %}y{% endif %}"),
        [],
        "ok",
    ),
    # false-positive guards
    Case("raw_ok", lambda d: d.add_paragraph("{% raw %}{% if x %}{% endraw %}"), [], "ok"),
    Case("set_block_ok", lambda d: d.add_paragraph("{% set y %}hi{% endset %}{{ y }}"), [], "ok"),
    Case("lt_ok", lambda d: d.add_paragraph("{% if v < 5 %}small{% endif %}"), [], "ok"),
    Case("expr_prefix_ok", lambda d: d.add_paragraph("{{p v }}"), [], "ok"),
    Case("macro_ok", lambda d: d.add_paragraph("{% macro m(a) %}[{{ a }}]{% endmacro %}{{ m(v) }}"), [], "ok"),
    Case("cross_para_warn", lambda d: [d.add_paragraph("A {{ v"), d.add_paragraph("}} B")], ["DTL002"], "ok"),
]


def build(case: Case, path) -> None:
    d = Document()
    case.build(d)
    d.save(path)
