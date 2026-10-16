"""v1.4: what docxtpl really sees in tracked changes and field codes (validated by rendering)."""

import re
import zipfile

from docx import Document
from docxtpl import DocxTemplate

from docxtpl_lint import lint_path
from docxtpl_lint.contract import infer_tree, to_schema

W_PLACE = "PLACEHOLDER"


def _mk(tmp_path, body_xml, name="t.docx"):
    d = Document()
    d.add_paragraph(W_PLACE)
    base = tmp_path / "base.docx"
    d.save(base)
    out = tmp_path / name
    with zipfile.ZipFile(base) as zin, zipfile.ZipFile(out, "w") as zout:
        for i in zin.infolist():
            data = zin.read(i.filename)
            if i.filename == "word/document.xml":
                s = data.decode()
                s = re.sub(r"<w:p>(?:(?!<w:p>).)*?PLACEHOLDER.*?</w:p>", body_xml, s, flags=re.S)
                data = s.encode()
            zout.writestr(i, data)
    return out


def _render_xml(path, ctx):
    t = DocxTemplate(str(path))
    t.render(ctx)
    out = path.with_name("out_" + path.name)
    t.save(out)
    return zipfile.ZipFile(out).read("word/document.xml").decode()


DEL_TAG = (
    '<w:p><w:r><w:t>A</w:t></w:r><w:del w:id="1" w:author="a"><w:r><w:delText>{{ old }}</w:delText></w:r></w:del>'
    "<w:r><w:t>B</w:t></w:r></w:p>"
)
DEL_INSIDE_TAG = (
    '<w:p><w:r><w:t xml:space="preserve">{{ </w:t></w:r><w:del w:id="1" w:author="a"><w:r><w:delText>nmae'
    '</w:delText></w:r></w:del><w:ins w:id="2" w:author="a"><w:r><w:t>name</w:t></w:r></w:ins><w:r>'
    '<w:t xml:space="preserve"> }}</w:t></w:r></w:p>'
)
HYPERLINK = (
    '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> HYPERLINK '
    '"{{ url }}" </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>link</w:t></w:r>'
    '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
)
FLDSIMPLE = (
    '<w:p><w:fldSimple w:instr=" HYPERLINK &quot;{{ site.url }}&quot; "><w:r><w:t>x</w:t></w:r></w:fldSimple></w:p>'
)
BAD_FIELD = '<w:p><w:r><w:instrText xml:space="preserve"> HYPERLINK "{% if x %}" </w:instrText></w:r></w:p>'


def test_deleted_code_is_rendered_and_reported(tmp_path):
    p = _mk(tmp_path, DEL_TAG)
    assert "<w:delText>SECRET</w:delText>" in _render_xml(p, {"old": "SECRET"})
    iss = lint_path(p).issues
    assert [i.rule for i in iss] == ["DTL003"] and iss[0].snippet == "{{ old }}"
    assert lint_path(p).variables == ["old"]  # docxtpl needs it, so it is part of the contract


def test_deletion_inside_a_tag_is_stripped_like_docxtpl(tmp_path):
    p = _mk(tmp_path, DEL_INSIDE_TAG)
    xml = _render_xml(p, {"name": "N", "nmae": "BAD", "nmaename": "WORSE"})
    assert ">N<" in xml and "BAD" not in xml and "WORSE" not in xml
    rep = lint_path(p)
    assert rep.issues == [] and rep.variables == ["name"]


def test_field_code_is_rendered_and_linted(tmp_path):
    p = _mk(tmp_path, HYPERLINK)
    assert 'HYPERLINK "http://x"' in _render_xml(p, {"url": "http://x"})
    assert lint_path(p).variables == ["url"]
    q = _mk(tmp_path, FLDSIMPLE, name="s.docx")
    assert "http://y" in _render_xml(q, {"site": {"url": "http://y"}})
    assert to_schema(infer_tree([q]))["properties"]["site"]["properties"] == {"url": {}}


def test_broken_field_code_fails_render_and_is_reported(tmp_path):
    p = _mk(tmp_path, BAD_FIELD)
    try:
        _render_xml(p, {"x": 1})
        raised = False
    except Exception:
        raised = True
    assert raised
    assert "DTL101" in [i.rule for i in lint_path(p).issues]


def test_plain_paragraphs_unchanged(tmp_path):
    d = Document()
    d.add_paragraph("Hello {{ name }} & <co>")
    d.save(tmp_path / "plain.docx")
    rep = lint_path(tmp_path / "plain.docx")
    assert rep.paragraphs["body"][0].text == "Hello {{ name }} & <co>"
