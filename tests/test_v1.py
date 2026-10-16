"""v1.0 features: infer / --schema / diff / verify / fix — each checked against real docxtpl."""

import json

import pytest
from docx import Document
from rule_cases import CASES, CTX, _table
from rule_cases import build as build_case

from docxtpl_lint.core import lint_path
from docxtpl_lint.cli import main
from docxtpl_lint.contract import diff, infer_tree, sample, to_schema
from docxtpl_lint.fixes import closest, suggest

docxtpl = pytest.importorskip("docxtpl")


def _doc(tmp_path, build, name="t.docx"):
    d = Document()
    build(d)
    p = tmp_path / name
    d.save(p)
    return p


def _render(path, ctx):
    tpl = docxtpl.DocxTemplate(str(path))
    tpl.render(ctx)
    out = path.with_suffix(".out.docx")
    tpl.save(str(out))
    d = Document(str(out))
    return [p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells]


INVOICE = [
    "Dear {{ customer.name }},",
    "{% for it in items %}{{ it.name }}: {{ it.price|round(2) }}{% endfor %}",
    "{% if total > 100 %}discount{% endif %}",
    "{% for k, v in extra.items() %}{{ k }}={{ v }}{% endfor %}",
    "{{ note|default('') }} {% if vat is defined %}{{ vat }}{% endif %}",
    "{% set even = count % 2 == 0 %}{{ even }}",
]


def _invoice(d):
    for t in INVOICE:
        d.add_paragraph(t)


# ---------------------------------------------------------------- infer


def test_infer_types(tmp_path):
    s = to_schema(infer_tree([_doc(tmp_path, _invoice)]))
    props = s["properties"]
    assert s["$schema"].endswith("2020-12/schema")
    assert props["customer"]["type"] == "object" and props["customer"]["required"] == ["name"]
    assert props["items"]["type"] == "array"
    assert props["items"]["items"]["properties"]["price"]["type"] == "number"  # |round
    assert "type" not in props["items"]["items"]["properties"]["name"]  # printed: any type
    assert props["total"]["type"] == "number"  # compared with a number
    assert props["count"]["type"] == "number"  # arithmetic
    assert props["extra"]["type"] == "object"  # .items()
    assert "note" not in s["required"] and "vat" not in s["required"]  # guarded
    assert "customer" in s["required"]


def test_infer_sample_renders_with_docxtpl(tmp_path):
    p = _doc(tmp_path, _invoice)
    text = "\n".join(_render(p, sample(infer_tree([p]))))
    assert "[customer.name]" in text and "[items[].name]" in text


@pytest.mark.parametrize("case", [c for c in CASES if c.docxtpl == "ok"], ids=lambda c: c.name)
def test_infer_sample_renders_every_ok_case(case, tmp_path):
    p = tmp_path / "c.docx"
    build_case(case, p)
    _render(p, sample(infer_tree([p])))


def test_infer_merges_templates(tmp_path):
    a = _doc(tmp_path, lambda d: d.add_paragraph("{{ name }} {{ a_only }}"), "a.docx")
    b = _doc(tmp_path, lambda d: d.add_paragraph("{{ name }} {% if b is defined %}{{ b }}{% endif %}"), "b.docx")
    s = to_schema(infer_tree([a, b]))
    assert s["properties"]["name"]["x-docxtpl-templates"] == ["a.docx", "b.docx"]
    assert s["properties"]["a_only"]["x-docxtpl-templates"] == ["a.docx"]
    assert set(s["required"]) == {"name", "a_only"}


def test_infer_cli(tmp_path, capsys):
    p = _doc(tmp_path, _invoice)
    assert main(["infer", str(p), "-o", str(tmp_path / "s.json"), "--sample", str(tmp_path / "d.json")]) == 0
    schema = json.loads((tmp_path / "s.json").read_text())
    data = json.loads((tmp_path / "d.json").read_text())
    assert len(data["items"]) == 2 and schema["type"] == "object"
    # the inferred contract is self-consistent
    assert main([str(p), "--schema", str(tmp_path / "s.json"), "--no-config"]) == 0
    assert main([str(p), "--data", str(tmp_path / "d.json"), "--no-config"]) == 0


# ---------------------------------------------------------------- --schema


def test_schema_check(tmp_path):
    p = _doc(
        tmp_path, lambda d: [d.add_paragraph("{{ name }} {{ typo }}"), d.add_paragraph("{% for x in n %}{% endfor %}")]
    )
    schema = {"type": "object", "properties": {"name": {}, "n": {"type": "number"}}}
    rep = lint_path(p, schema=schema)
    assert [i.rule for i in rep.issues] == ["DTL304", "DTL305"]
    assert rep.issues[0].params["path"] == "typo"


# ---------------------------------------------------------------- diff


def test_contract_diff(tmp_path):
    old = _doc(
        tmp_path,
        lambda d: d.add_paragraph("{{ a }} {{ b }} {{ c.x }} {% if o is defined %}{{ o }}{% endif %}"),
        "o.docx",
    )
    new = _doc(
        tmp_path,
        lambda d: d.add_paragraph("{{ a }} {{ bb }} {% for i in c %}{{ i }}{% endfor %} {{ o }}"),
        "n.docx",
    )
    ch = {(c.kind, c.path): c for c in diff(infer_tree([old]), infer_tree([new]))}
    assert ch[("added", "bb")].breaking and ch[("added", "bb")].detail == "renamed from b?"
    assert not ch[("removed", "b")].breaking
    assert ch[("type", "c")].breaking and ch[("type", "c")].detail == "object -> array"
    assert ch[("optional", "o")].breaking
    assert ("added", "a") not in ch
    assert main(["diff", str(old), str(new), "--fail-on-breaking"]) == 1
    assert main(["diff", str(old), str(old), "--fail-on-breaking"]) == 0


# ---------------------------------------------------------------- verify (render equivalence)

LOSSY = {
    # docxtpl replaces the paragraph with the {%p %} tag: `{{ total }}` disappears, no error
    "p_same_paragraph": lambda d: [d.add_paragraph("Total {{ total }} {%p if x %}"), d.add_paragraph("{%p endif %}")],
    # docxtpl replaces the row with the {%tr %} tag: `{{ i.name }}` disappears, no error
    "tr_same_cell": lambda d: _table(d, 2, 1, {(0, 0): "{%tr for i in xs %} {{ i.name }}", (1, 0): "{%tr endfor %}"}),
}


@pytest.mark.parametrize("name", LOSSY)
def test_verify_finds_silently_lost_output(name, tmp_path, capsys):
    p = _doc(tmp_path, LOSSY[name])
    assert main(["verify", str(p)]) == 1
    assert "DTL401" in capsys.readouterr().out


def test_verify_clean_and_conditional(tmp_path, capsys):
    p = _doc(tmp_path, _invoice)
    assert main(["verify", str(p)]) == 0
    # inside a false branch: skipped, not a false positive
    q = _doc(tmp_path, lambda d: d.add_paragraph("{% if not flag %}{{ hidden }}{% endif %}"), "q.docx")
    assert main(["verify", str(q)]) == 0


def test_verify_render_failure(tmp_path, capsys):
    p = _doc(tmp_path, lambda d: d.add_paragraph("{{ fmt(x) }}"))
    assert main(["verify", str(p), "--format", "json"]) == 1
    out = json.loads(capsys.readouterr().out)
    assert out["files"][0]["issues"][0]["rule"] == "DTL402"


# ---------------------------------------------------------------- fix suggestions


def _apply(path, fixes):
    d = Document(str(path))
    paras = list(d.paragraphs) + [p for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs]
    for fx in fixes:
        for p in paras:
            if p.text == fx.old:
                for r in p.runs[1:]:
                    r._r.getparent().remove(r._r)
                p.runs[0].text = fx.new
                break
    out = path.with_name("fixed.docx")
    d.save(str(out))
    return out


FIXABLE = {c.name: c.build for c in CASES if c.rules and c.rules[0] in ("DTL202", "DTL203", "DTL204")}
FIXABLE["typo_endfro"] = lambda d: [d.add_paragraph("{% for i in xs %}{{ i }}"), d.add_paragraph("{% endfro %}")]
FIXABLE["stray_endif"] = lambda d: [d.add_paragraph("{{ v }}"), d.add_paragraph("{% endif %}")]
FIXABLE["unknown_tag"] = lambda d: [d.add_paragraph("{% iff x %}yes{% endif %}")]


@pytest.mark.parametrize("name", FIXABLE)
def test_fix_suggestion_renders_with_docxtpl(name, tmp_path):
    p = _doc(tmp_path, FIXABLE[name])
    rep = lint_path(p)
    assert rep.issues
    fixes = suggest(rep)
    assert fixes, [i.rule for i in rep.issues]
    fixed = _apply(p, fixes)
    assert lint_path(fixed).issues == []
    text = _render(fixed, CTX)
    if name == "dtl204_bad":  # the loop really produces one row per item now
        assert text == ["1", "2"]


def test_fix_did_you_mean_data(tmp_path, capsys):
    p = _doc(tmp_path, lambda d: d.add_paragraph("{{ custmer.name }}"))
    (tmp_path / "d.json").write_text(json.dumps({"customer": {"name": "A"}}))
    assert main(["fix", str(p), "--data", str(tmp_path / "d.json")]) == 0
    out = capsys.readouterr().out
    assert "-{{ custmer.name }}" in out and "+{{ customer.name }}" in out
    assert _render(
        _apply(p, suggest(lint_path(p, {"customer": {"name": "A"}}), [{"customer": {"name": "A"}}])),
        {"customer": {"name": "A"}},
    ) == ["A"]


def test_closest_is_conservative():
    assert closest("endfro", ["endfor", "endif"]) == "endfor"
    assert closest("banana", ["endfor", "endif"]) is None  # too far
    assert closest("ab", ["ac", "ad"]) is None  # tie: no guess


def test_verify_same_path_twice_one_branch_hidden(tmp_path):
    # DeepSeek round 9 claimed a false DTL401 here; one sentinel per path, so no.
    p = _doc(tmp_path, lambda d: [d.add_paragraph("{{ x }}"), d.add_paragraph("{%p if flag is undefined %}"),
                                  d.add_paragraph("{{ x }}"), d.add_paragraph("{%p endif %}")])  # fmt: skip
    assert main(["verify", str(p)]) == 0
