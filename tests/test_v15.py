"""v1.5: deeper data contract: list filters, attribute filters, RichText/Subdoc tags (render-validated)."""

import pytest
from docx import Document
from docxtpl import DocxTemplate, RichText

from docxtpl_lint.contract import diff, infer_tree, sample, to_pydantic, to_schema
from docxtpl_lint.core import lint_path


def _doc(tmp_path, *paras, name="t.docx"):
    d = Document()
    for p in paras:
        d.add_paragraph(p)
    path = tmp_path / name
    d.save(path)
    return path


def _render(path, ctx):
    t = DocxTemplate(str(path))
    t.render(ctx)
    out = path.with_name("out_" + path.name)
    t.save(out)
    return [p.text for p in Document(out).paragraphs]


def _rules(p, data):
    return [i.rule for i in lint_path(p, data).issues]


def test_join_on_string_splits_silently_and_number_raises(tmp_path):
    p = _doc(tmp_path, "{{ tags|join(', ') }}")
    assert _render(p, {"tags": "abc"}) == ["a, b, c"]
    with pytest.raises(TypeError):
        _render(p, {"tags": 5})
    assert _rules(p, {"tags": "abc"}) == ["DTL306"]
    assert _rules(p, {"tags": 5}) == ["DTL306"]
    assert _rules(p, {"tags": ["a", "b"]}) == []
    # DeepSeek round 14: a mapping is iterated by its keys, silently
    assert _render(p, {"tags": {"x": 1, "y": 2}}) == ["x, y"]
    assert _rules(p, {"tags": {"x": 1}}) == ["DTL306"]
    assert to_schema(infer_tree([p]))["properties"]["tags"]["type"] == "array"


def test_sum_attribute_infers_item_number(tmp_path):
    p = _doc(tmp_path, "{{ items|sum(attribute='price') }}")
    assert _render(p, {"items": [{"price": 1}, {"price": 2}]}) == ["3"]
    with pytest.raises(TypeError):
        _render(p, {"items": [{"price": "1"}]})
    s = to_schema(infer_tree([p]))
    assert s["properties"]["items"]["items"]["properties"]["price"]["type"] == "number"
    # the inferred sample renders
    assert _render(p, sample(infer_tree([p]))) == ["2"]


def test_loop_over_filtered_list_keeps_item_fields(tmp_path):
    p = _doc(tmp_path, "{% for i in items|selectattr('ok')|sort(attribute='n') %}{{ i.title }}{% endfor %}")
    tree = infer_tree([p])
    props = to_schema(tree)["properties"]["items"]["items"]["properties"]
    assert set(props) == {"ok", "n", "title"}
    data = {
        "items": [{"ok": 1, "n": 2, "title": "B"}, {"ok": 1, "n": 1, "title": "A"}, {"ok": 0, "n": 0, "title": "X"}]
    }
    assert _render(p, data) == ["AB"]
    assert _rules(p, {"items": [{"ok": 1, "n": 1}]}) == ["DTL301"]  # items[*].title missing


def test_selectattr_first_argument_is_an_attribute(tmp_path):
    """DeepSeek claimed selectattr('odd') names a test; Jinja's signature is selectattr(attribute, test)."""
    p = _doc(tmp_path, "{% for i in xs|selectattr('odd') %}{{ i.v }}{% endfor %}")
    assert _render(p, {"xs": [{"odd": 1, "v": "a"}, {"odd": 0, "v": "b"}]}) == ["a"]
    assert set(to_schema(infer_tree([p]))["properties"]["xs"]["items"]["properties"]) == {"odd", "v"}


def test_map_attribute_join(tmp_path):
    p = _doc(tmp_path, "{{ people|map(attribute='name')|join(', ') }}")
    assert _render(p, {"people": [{"name": "a"}, {"name": "b"}]}) == ["a, b"]
    assert "name" in to_schema(infer_tree([p]))["properties"]["people"]["items"]["properties"]


def test_first_on_string_is_not_flagged(tmp_path):
    p = _doc(tmp_path, "{{ name|first }}")
    assert _render(p, {"name": "Ada"}) == ["A"]
    assert _rules(p, {"name": "Ada"}) == []


def test_richtext_plain_value_renders_empty(tmp_path):
    p = _doc(tmp_path, "A {{r note }} B")
    assert _render(p, {"note": "plain"}) == ["A  B"]
    assert _render(p, {"note": RichText("rich")}) == ["A rich B"]
    assert _rules(p, {"note": "plain"}) == ["DTL307"]
    assert _render(p, {"note": {"text": "x"}}) == ["A  B"]  # a dict is no better
    assert _rules(p, {"note": {"text": "x"}}) == ["DTL307"]
    s = to_schema(infer_tree([p]))["properties"]["note"]
    assert s["x-docxtpl-object"] == "RichText"
    assert "docxtpl.RichText" in to_pydantic(infer_tree([p]))


def test_diff_plain_to_richtext_is_breaking(tmp_path):
    a = _doc(tmp_path, "{{ note }}", name="a.docx")
    b = _doc(tmp_path, "{{r note }}", name="b.docx")
    ch = diff(infer_tree([a]), infer_tree([b]))
    assert [(c.kind, c.breaking, c.detail) for c in ch] == [("type", True, "any -> richtext")]
