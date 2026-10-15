"""Build the .docx fixtures used by the tests and README (python-docx only).

Each broken template models a real docxtpl issue:
  * #603 – loops inside tables  ->  issue603_table_loop.docx
  * #644 – `%` / `}` inside {%p %} tags  ->  issue644_percent.docx
  * #378 – "where in the template is the error?"  ->  issue378_unclosed.docx
Run:  python tests/make_fixtures.py [outdir]
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document


def _split_runs(paragraph, *chunks: str) -> None:
    """Write text as several runs, like Word does after spell-check / formatting edits."""
    for i, ch in enumerate(chunks):
        r = paragraph.add_run(ch)
        r.bold = i % 2 == 1


def clean(out: Path) -> None:
    d = Document()
    p = d.add_paragraph("Dear ")
    _split_runs(p, "{{ cust", "omer.name", " }},")  # split across runs: fine for docxtpl
    d.add_paragraph("{%p if items %}")
    t = d.add_table(rows=4, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Price"
    t.cell(1, 0).text = "{%tr for item in items %}"
    t.cell(2, 0).text, t.cell(2, 1).text = "{{ item.name }}", "{{ item.price }}"
    t.cell(3, 0).text = "{%tr endfor %}"
    d.add_paragraph("{%p endif %}")
    d.add_paragraph("{% set total = items | sum(attribute='price') %}Total: {{ total }}")
    d.sections[0].footer.paragraphs[0].text = "{{ company }}"
    d.save(out / "clean.docx")


def issue603(out: Path) -> None:
    d = Document()
    d.add_paragraph("Order {{ order.id }}")
    # table 1: tr open/close in the SAME row -> row eaten, loop broken
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Name", "Qty"
    t.cell(1, 0).text = "{%tr for line in order.lines %}{{ line.name }}"
    t.cell(1, 1).text = "{{ line.qty }}{%tr endfor %}"
    # table 2: opened with plain {% for %}, closed with {%tr endfor %}
    t2 = d.add_table(rows=3, cols=1)
    t2.cell(0, 0).text = "{% for c in order.comments %}"
    t2.cell(1, 0).text = "{{ c }}"
    t2.cell(2, 0).text = "{%tr endfor %}"
    # table 3: space before prefix
    t3 = d.add_table(rows=3, cols=1)
    t3.cell(0, 0).text = "{% tr for x in order.extra %}"
    t3.cell(1, 0).text = "{{ x }}"
    t3.cell(2, 0).text = "{%tr endfor %}"
    d.save(out / "issue603_table_loop.docx")


def issue644(out: Path) -> None:
    d = Document()
    d.add_paragraph("{%p if row_index % 2 == 0 %}")
    d.add_paragraph("even row")
    d.add_paragraph("{%p endif %}")
    d.add_paragraph("{%p set label = {'a': 1}['a'] %}")
    d.save(out / "issue644_percent.docx")


def issue378(out: Path) -> None:
    d = Document()
    d.add_paragraph("Invoice for {{ client.name }}")
    d.add_table(rows=1, cols=1).cell(0, 0).text = "header"
    t = d.add_table(rows=3, cols=2)
    t.cell(0, 0).text = "Status"
    t.cell(2, 0).text = "{% if paid %}PAID"  # never closed
    t.cell(2, 1).text = "{{ amount }}"
    d.add_paragraph("Contact: {{ client.")
    d.add_paragraph("email }}")
    d.add_paragraph("{% endfor %}")  # stray
    d.save(out / "issue378_unclosed.docx")


def data(out: Path) -> None:
    (out / "clean.data.json").write_text(
        json.dumps(
            {"customer": {"name": "ACME"}, "items": [{"name": "Pen"}], "company": "Foo Ltd", "unused_key": 1},
            indent=2,
        ),
        encoding="utf-8",
    )


def build(out: Path) -> Path:
    out.mkdir(parents=True, exist_ok=True)
    for fn in (clean, issue603, issue644, issue378, data):
        fn(out)
    return out


if __name__ == "__main__":
    print(build(Path(sys.argv[1] if len(sys.argv) > 1 else "examples")))
