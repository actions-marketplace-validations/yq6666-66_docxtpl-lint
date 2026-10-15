"""Reader tests (first draft by DeepSeek, reviewed and kept as-is)."""

from docx import Document
from lxml import etree

from docxtpl_lint.reader import iter_paragraphs

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_W = f"{{{W}}}"


def _save_and_reload(doc, tmp_path, mutate):
    """Save doc, reopen the document.xml part, apply mutate(root), write back, return path."""
    path = tmp_path / "t.docx"
    doc.save(str(path))
    import zipfile

    with zipfile.ZipFile(str(path)) as zf:
        parts = {n: zf.read(n) for n in zf.namelist()}
    root = etree.fromstring(parts["word/document.xml"])
    mutate(root)
    parts["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(str(path), "w", zipfile.ZIP_DEFLATED) as zf:
        for n, data in parts.items():
            zf.writestr(n, data)
    return path


def _first_body_p(root):
    return root.find(f"{_W}body").find(f"{_W}p")


def test_ins_counted_del_excluded(tmp_path):
    doc = Document()
    doc.add_paragraph("keep")

    def mutate(root):
        p = _first_body_p(root)
        ins = etree.SubElement(p, f"{_W}ins")
        ins.set(f"{_W}id", "1")
        ins.set(f"{_W}author", "a")
        ins.set(f"{_W}date", "2020-01-01T00:00:00Z")
        r_ins = etree.SubElement(ins, f"{_W}r")
        t_ins = etree.SubElement(r_ins, f"{_W}t")
        t_ins.text = "INSERTED"

        dele = etree.SubElement(p, f"{_W}del")
        dele.set(f"{_W}id", "2")
        dele.set(f"{_W}author", "a")
        dele.set(f"{_W}date", "2020-01-01T00:00:00Z")
        r_del = etree.SubElement(dele, f"{_W}r")
        dt = etree.SubElement(r_del, f"{_W}delText")
        dt.text = "DELETED"

    path = _save_and_reload(doc, tmp_path, mutate)
    paras = [p for p in iter_paragraphs(path) if p.part == "body"]
    text = paras[0].text
    assert "INSERTED" in text
    assert "DELETED" not in text


def test_textbox_paragraph_is_separate(tmp_path):
    doc = Document()
    doc.add_paragraph("outer")

    def mutate(root):
        p = _first_body_p(root)
        r = etree.SubElement(p, f"{_W}r")
        t = etree.SubElement(r, f"{_W}t")
        t.text = "OUTER"
        pict = etree.SubElement(r, f"{_W}pict")
        txbx = etree.SubElement(pict, f"{_W}txbxContent")
        inner_p = etree.SubElement(txbx, f"{_W}p")
        inner_r = etree.SubElement(inner_p, f"{_W}r")
        inner_t = etree.SubElement(inner_r, f"{_W}t")
        inner_t.text = "INNERBOX"

    path = _save_and_reload(doc, tmp_path, mutate)
    paras = [p for p in iter_paragraphs(path) if p.part == "body"]
    texts = [p.text for p in paras]
    assert "INNERBOX" in texts
    outer = next(p for p in paras if "OUTER" in p.text)
    assert "INNERBOX" not in outer.text
    inner = next(p for p in paras if "INNERBOX" in p.text)
    assert "OUTER" not in inner.text


def test_nested_table_cells(tmp_path):
    doc = Document()
    doc.add_paragraph("top")

    def mutate(root):
        body = root.find(f"{_W}body")
        tbl = etree.SubElement(body, f"{_W}tbl")
        tr = etree.SubElement(tbl, f"{_W}tr")
        tc = etree.SubElement(tr, f"{_W}tc")
        p = etree.SubElement(tc, f"{_W}p")
        r = etree.SubElement(p, f"{_W}r")
        t = etree.SubElement(r, f"{_W}t")
        t.text = "OUTERCELL"

        inner_tbl = etree.SubElement(tc, f"{_W}tbl")
        inner_tr = etree.SubElement(inner_tbl, f"{_W}tr")
        inner_tc = etree.SubElement(inner_tr, f"{_W}tc")
        inner_p = etree.SubElement(inner_tc, f"{_W}p")
        inner_r = etree.SubElement(inner_p, f"{_W}r")
        inner_t = etree.SubElement(inner_r, f"{_W}t")
        inner_t.text = "INNERCELL"

    path = _save_and_reload(doc, tmp_path, mutate)
    paras = [p for p in iter_paragraphs(path) if p.part == "body"]

    outer = next(p for p in paras if p.text == "OUTERCELL")
    assert len(outer.cells) == 1
    assert (outer.cells[0].table, outer.cells[0].row, outer.cells[0].col) == (1, 1, 1)
    assert outer.cell_para == 1

    inner = next(p for p in paras if p.text == "INNERCELL")
    assert len(inner.cells) == 2
    assert (inner.cells[0].table, inner.cells[0].row, inner.cells[0].col) == (1, 1, 1)
    assert (inner.cells[1].table, inner.cells[1].row, inner.cells[1].col) == (2, 1, 1)
    assert inner.cell_para == 1
