"""v1.7: fix --write applies fixes to a copy, inside the existing runs (render-validated)."""

import pytest
from docx import Document
from docxtpl import DocxTemplate
from jinja2 import TemplateSyntaxError

from docxtpl_lint.cli import main
from docxtpl_lint.core import lint_path
from docxtpl_lint.fixes import suggest, write_fixed


def _render(path, ctx):
    t = DocxTemplate(str(path))
    t.render(ctx)
    out = path.with_name("out_" + path.name)
    t.save(out)
    d = Document(out)
    return [p.text for p in d.paragraphs], d


def _broken(tmp_path):
    d = Document()
    p = d.add_paragraph()
    r = p.add_run("Total: ")
    r.bold = True
    p.add_run("{{ to")
    p.add_run("tal }")  # tag split over runs, missing a brace
    d.add_paragraph("{% for x in xs %}{{ x }}{% endfro %}")
    t = d.add_table(rows=3, cols=1)
    t.cell(0, 0).text = "{% tr for r in rows %}"
    t.cell(1, 0).text = "{{ r }}"
    t.cell(2, 0).text = "{%tr endfor %}"
    path = tmp_path / "inv.docx"
    d.save(path)
    return path


def test_write_fixes_copy_keeps_formatting_and_renders(tmp_path):
    src = _broken(tmp_path)
    before = src.read_bytes()
    with pytest.raises(TemplateSyntaxError):
        _render(src, {"total": 5, "xs": ["a"], "rows": ["r1"]})
    out = tmp_path / "inv.fixed.docx"
    applied = write_fixed(str(src), str(out))
    assert sorted(f.issue.rule for f in applied) == ["DTL001", "DTL102", "DTL202"]
    assert src.read_bytes() == before  # the template is untouched
    assert lint_path(out).issues == []
    texts, doc = _render(out, {"total": 5, "xs": ["a", "b"], "rows": ["r1", "r2"]})
    assert texts[:2] == ["Total: 5", "ab"]
    assert [c.text for c in doc.tables[0].columns[0].cells] == ["r1", "r2"]
    runs = Document(out).paragraphs[0].runs
    assert runs[0].text == "Total: " and runs[0].bold  # formatting of untouched runs kept


def test_unsafe_fixes_need_flag(tmp_path):
    d = Document()
    d.add_paragraph("{{ custmer.name }}")
    p = tmp_path / "t.docx"
    d.save(p)
    ctx = [{"customer": {"name": "Ada"}}]
    fx = suggest(lint_path(p, contexts=ctx), ctx)
    assert [(f.issue.rule, f.safe) for f in fx] == [("DTL301", False)]
    assert write_fixed(str(p), str(tmp_path / "a.docx"), ctx) == []
    assert len(write_fixed(str(p), str(tmp_path / "b.docx"), ctx, unsafe=True)) == 1
    assert _render(tmp_path / "b.docx", ctx[0])[0] == ["Ada"]


def test_refuses_to_overwrite_template(tmp_path):
    src = _broken(tmp_path)
    with pytest.raises(ValueError):
        write_fixed(str(src), str(src))


def test_tracked_change_paragraphs_left_alone(tmp_path):
    from test_v14 import _mk

    p = _mk(tmp_path, '<w:p><w:r><w:t>{% endfro %}</w:t></w:r><w:del w:id="1" w:author="a"><w:r><w:delText>x'
                      "</w:delText></w:r></w:del></w:p>")  # fmt: skip
    assert write_fixed(str(p), str(tmp_path / "o.docx")) == []


def test_cli_write_dir_and_diff_marks_unsafe(tmp_path, capsys):
    a = _broken(tmp_path)
    assert main(["fix", str(a), str(a), "--write", str(tmp_path / "fixed")]) == 0
    assert (tmp_path / "fixed" / "inv.fixed.docx").exists()
    d = Document()
    d.add_paragraph("{% endfor %}")
    q = tmp_path / "q.docx"
    d.save(q)
    capsys.readouterr()
    assert main(["fix", str(q)]) == 0
    assert "unsafe" in capsys.readouterr().out
