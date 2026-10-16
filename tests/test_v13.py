"""v1.3: Jinja correctness (DTL104/105/106), each validated with a real docxtpl render."""

import sys

import pytest
from docx import Document
from docxtpl import DocxTemplate
from jinja2 import Environment, TemplateAssertionError, TemplateRuntimeError, UndefinedError

from docxtpl_lint.core import lint_path
from docxtpl_lint.cli import main


def _doc(tmp_path, *paras, name="t.docx"):
    d = Document()
    for p in paras:
        d.add_paragraph(p)
    path = tmp_path / name
    d.save(path)
    return path


def _render(path, ctx, env=None):
    t = DocxTemplate(str(path))
    t.render(ctx, jinja_env=env)
    out = path.with_name("out_" + path.name)
    t.save(out)
    return [p.text for p in Document(out).paragraphs]


def _rules(p, **kw):
    return [i.rule for i in lint_path(p, **kw).issues]


def test_dtl104_unknown_filter_fails_in_docxtpl(tmp_path):
    p = _doc(tmp_path, "{{ total|money }}")
    with pytest.raises(TemplateAssertionError):
        _render(p, {"total": 1})
    assert _rules(p) == ["DTL104"]
    env = Environment()
    env.filters["money"] = lambda v: f"${v:.2f}"
    assert _render(p, {"total": 1}, env) == ["$1.00"]
    assert _rules(p, filters={"money"}) == []


def test_dtl104_filter_block_and_builtin(tmp_path):
    p = _doc(tmp_path, "{% filter shout %}x{% endfilter %} {{ a|upper|default('z') }}")
    assert _rules(p) == ["DTL104"]


def test_dtl105_unknown_test_fails_in_docxtpl(tmp_path):
    p = _doc(tmp_path, "{% if x is money %}a{% endif %}{% if x is number %}b{% endif %}")
    with pytest.raises(TemplateRuntimeError):
        _render(p, {"x": 1})
    assert _rules(p) == ["DTL105"]
    assert _rules(p, tests={"money"}) == []


def test_dtl106_set_in_loop_renders_empty(tmp_path):
    p = _doc(tmp_path, "{% for i in items %}{% set last = i %}{% endfor %}[{{ last }}]")
    assert _render(p, {"items": [1, 2]}) == ["[]"]
    assert _rules(p) == ["DTL106"]


def test_dtl106_loop_target_attr_raises(tmp_path):
    p = _doc(tmp_path, "{% for i in items %}{% endfor %}", "{{ i.n }}")
    with pytest.raises(UndefinedError):
        _render(p, {"items": [{"n": 1}]})
    iss = lint_path(p).issues
    assert [i.rule for i in iss] == ["DTL106"] and iss[0].location.paragraph == 2
    assert "i" not in lint_path(p).variables  # not part of the data contract


def test_dtl106_loop_outside_loop(tmp_path):
    p = _doc(tmp_path, "{{ loop.index }}")
    assert _rules(p) == ["DTL106"]


def test_dtl106_no_false_positives(tmp_path):
    ok = [
        "{% for i in items %}{{ loop.index }}{{ i.n }}{% endfor %}",
        "{% set ns = namespace(t=0) %}{% for i in items %}{% set ns.t = ns.t + i.p %}{% endfor %}{{ ns.t }}",
        "{{ item }}{% for item in items %}{{ item }}{% endfor %}{{ item }}",  # also real top-level data
        "{% for a in xs %}{% for b in a.ys %}{{ b }}{% endfor %}{{ a }}{% endfor %}",
        "{% macro m(i) %}{{ i }}{% endmacro %}{% for i in xs %}{% endfor %}{{ m(1) }}",
        "{% for i in xs %}{% endfor %}{% if c %}{% set i = 5 %}{% endif %}{{ i }}",  # DeepSeek round 12
        "{% for i in xs %}{% endfor %}{% set i = 5 %}{{ i }}",
        "{% macro m(loop) %}{{ loop }}{% endmacro %}",
    ]
    for k, src in enumerate(ok):
        assert _rules(_doc(tmp_path, src, name=f"ok{k}.docx")) == [], src
    # the third one really renders the top-level value after the loop
    assert _render(_doc(tmp_path, ok[2], name="r.docx"), {"item": "T", "items": ["a"]}) == ["TaT"]


def test_cli_filters_config_and_jinja_env(tmp_path, monkeypatch, capsys):
    p = _doc(tmp_path, "{{ x|money }} {% if x is vat %}{% endif %}")
    assert main([str(p), "--no-config"]) == 1
    assert main([str(p), "--no-config", "--filters", "money", "--tests", "vat"]) == 0
    (tmp_path / "myenv.py").write_text(
        "from jinja2 import Environment\n"
        "def make():\n    e = Environment()\n    e.filters['money'] = str\n    e.tests['vat'] = bool\n    return e\n"
    )
    monkeypatch.chdir(tmp_path)
    assert main([str(p), "--no-config", "--jinja-env", "myenv:make"]) == 0
    sys.modules.pop("myenv", None)
    assert main([str(p), "--no-config", "--jinja-env", "myenv"]) == 2
    assert main([str(p), "--no-config", "--jinja-env", "nope_mod:x"]) == 2
    (tmp_path / "pyproject.toml").write_text('[tool.docxtpl-lint]\nfilters = ["money"]\ntests = ["vat"]\n')
    assert main([str(p)]) == 0


def test_broken_template_no_jinja_checks(tmp_path):
    p = _doc(tmp_path, "{% if a %}{{ x|money }}")
    assert _rules(p) == ["DTL101"]
