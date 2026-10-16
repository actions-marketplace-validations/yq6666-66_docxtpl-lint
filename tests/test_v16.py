"""v1.6: escaping (DTL308) and a stronger verify, validated with real docxtpl renders."""

import json

import pytest
from docx import Document
from docxtpl import DocxTemplate

from docxtpl_lint.cli import main
from docxtpl_lint.core import lint_path
from docxtpl_lint.verify import verify_path


def _doc(tmp_path, *paras, name="t.docx"):
    d = Document()
    for p in paras:
        d.add_paragraph(p)
    path = tmp_path / name
    d.save(path)
    return path


def _render(path, ctx, **kw):
    t = DocxTemplate(str(path))
    t.render(ctx, **kw)
    out = path.with_name("out_" + path.name)
    t.save(out)
    return [p.text for p in Document(out).paragraphs]


@pytest.mark.parametrize(
    ("value", "lost"),
    [("Smith & Co", "[Smith  Co]"), ("a < b", "[a  b]"), ("<b>bold</b>", "[")],
)
def test_unescaped_values_are_cut(tmp_path, value, lost):
    p = _doc(tmp_path, "[{{ c.name }}]")
    assert _render(p, {"c": {"name": value}}) == [lost]
    assert _render(p, {"c": {"name": value}}, autoescape=True) == [f"[{value}]"]
    assert [i.rule for i in lint_path(p, {"c": {"name": value}}).issues] == ["DTL308"]
    assert lint_path(p, {"c": {"name": value}}, autoescape=True).issues == []


def test_escaped_or_harmless_values_not_reported(tmp_path):
    p = _doc(tmp_path, "[{{ a|e }}] [{{ b }}] {% if c %}x{% endif %} {{r d }}")
    data = {"a": "A & B", "b": 'x > y "q"', "c": "C & D", "d": "x"}
    assert _render(p, {**data, "d": ""})[0].startswith('[A & B] [x > y "q"]')
    assert "DTL308" not in [i.rule for i in lint_path(p, data).issues]


def test_filtered_output_still_reported(tmp_path):
    p = _doc(tmp_path, "{{ n|upper }}")
    assert _render(p, {"n": "a & b"}) == ["A  B"]
    assert [i.rule for i in lint_path(p, {"n": "a & b"}).issues] == ["DTL308"]


def test_concat_and_condexpr_reported(tmp_path):
    """DeepSeek round 15: `~` and inline if were missed."""
    p = _doc(tmp_path, "{{ 'Co: ' ~ a }} {{ b if c else 'n/a' }}")
    data = {"a": "x & y", "b": "<i>", "c": 1}
    assert _render(p, data) == ["Co: x  y "]
    assert [i.rule for i in lint_path(p, data).issues] == ["DTL308", "DTL308"]


def test_cli_autoescape_flag_and_config(tmp_path, monkeypatch):
    p = _doc(tmp_path, "{{ n }}")
    (tmp_path / "d.json").write_text(json.dumps({"n": "a & b"}))
    assert main([str(p), "--data", str(tmp_path / "d.json"), "--no-config"]) == 1
    assert main([str(p), "--data", str(tmp_path / "d.json"), "--no-config", "--autoescape"]) == 0
    (tmp_path / "pyproject.toml").write_text("[tool.docxtpl-lint]\nautoescape = true\n")
    monkeypatch.chdir(tmp_path)
    assert main([str(p), "--data", str(tmp_path / "d.json")]) == 0


def test_verify_checks_richtext_outputs(tmp_path):
    p = _doc(tmp_path, "{{r x }} {%p if y %}", "in", "{%p endif %}")
    assert [i.rule for i in verify_path(p).issues] == ["DTL401"]
    ok = _doc(tmp_path, "A {{r x }} B", "{{p sd }}", name="ok.docx")
    assert verify_path(ok).issues == []


def test_verify_multiple_contexts_and_out(tmp_path, capsys):
    p = _doc(tmp_path, "{{ items|join(',') }}")
    (tmp_path / "a.json").write_text('{"items": ["x", "y"]}')
    (tmp_path / "b.json").write_text('{"items": 5}')
    rc = main(["verify", str(p), "--data", str(tmp_path / "a.json"), "--data", str(tmp_path / "b.json"),
               "--out", str(tmp_path / "out")])  # fmt: skip
    assert rc == 1
    assert "(--data #2)" in capsys.readouterr().out
    assert (tmp_path / "out" / "t.rendered.1.docx").exists()
    assert [x.text for x in Document(tmp_path / "out" / "t.rendered.1.docx").paragraphs] == ["x,y"]


def test_verify_autoescape_and_jinja_env(tmp_path, monkeypatch):
    p = _doc(tmp_path, "{{ v|money }}")
    assert [i.rule for i in verify_path(p).issues] == ["DTL402"]
    (tmp_path / "envmod.py").write_text(
        "from jinja2 import Environment\nenv = Environment()\nenv.filters['money'] = str\n"
    )
    monkeypatch.chdir(tmp_path)
    assert main(["verify", str(p), "--jinja-env", "envmod:env"]) == 0
    q = _doc(tmp_path, "{{ n }}", name="q.docx")
    (tmp_path / "d.json").write_text('{"n": "a & b"}')
    assert (
        main(["verify", str(q), "--data", str(tmp_path / "d.json"), "--autoescape", "--out", str(tmp_path / "o")]) == 0
    )
    assert [x.text for x in Document(tmp_path / "o" / "q.rendered.docx").paragraphs] == ["a & b"]
