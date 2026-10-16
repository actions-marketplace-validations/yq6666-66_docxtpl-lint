import json

import pytest
from docx import Document
from rule_cases import CASES
from rule_cases import build as build_case

from docxtpl_lint.core import lint_path
from docxtpl_lint.cli import main
from docxtpl_lint.messages import RULES
from docxtpl_lint.reader import CellPos


def rules(rep):
    return [i.rule for i in rep.issues]


def _doc(tmp_path, build, name="t.docx"):
    d = Document()
    build(d)
    path = tmp_path / name
    d.save(path)
    return path


# --- one case per rule, cross-checked against real docxtpl in test_docxtpl_behaviour.py


@pytest.mark.parametrize("case", CASES, ids=lambda c: c.name)
def test_rule_cases(case, tmp_path):
    p = tmp_path / "c.docx"
    build_case(case, p)
    assert rules(lint_path(p)) == case.rules


def test_every_rule_has_bilingual_message():
    for r in RULES.values():
        assert r.en and r.zh and r.severity in ("error", "warning")


# --- fixtures modelled on upstream issues -------------------------------------


def test_clean_template_has_no_issues(fx):
    rep = lint_path(fx / "clean.docx")
    assert rep.issues == []
    assert "items[].price" in rep.variables
    assert "total" not in rep.variables  # {% set %} local
    assert "customer.name" in rep.variables  # tag split over three runs


def test_378_unclosed_if_located_to_cell(fx):
    rep = lint_path(fx / "issue378_unclosed.docx")
    (dtl101,) = [i for i in rep.issues if i.rule == "DTL101"]
    assert dtl101.location.cells == (CellPos(table=2, row=3, col=1),)
    assert dtl101.snippet == "{% if paid %}"
    assert "T2 R3 C1" in dtl101.location.human()
    assert "表2 行3 列1" in dtl101.location.human("zh")


def test_378_reports_all_problems(fx):
    assert sorted(rules(lint_path(fx / "issue378_unclosed.docx"))) == ["DTL002", "DTL101", "DTL102"]


def test_603_rules(fx):
    assert rules(lint_path(fx / "issue603_table_loop.docx")) == ["DTL206", "DTL204", "DTL202"]


def test_644_percent_and_brace(fx):
    rep = lint_path(fx / "issue644_percent.docx")
    assert rules(rep) == ["DTL201", "DTL201"]
    assert [i.location.paragraph for i in rep.issues] == [1, 4]


# --- tag scanning ----------------------------------------------------------------


def test_never_closed_is_error(tmp_path):
    p = _doc(tmp_path, lambda d: [d.add_paragraph("{{ a"), d.add_paragraph("b")])
    (i,) = lint_path(p).issues
    assert (i.rule, i.severity, i.location.paragraph) == ("DTL001", "error", 1)


def test_cross_paragraph_opt_in(fx):
    loose = lint_path(fx / "issue378_unclosed.docx", allow_cross_paragraph=True)
    assert "DTL002" not in rules(loose)
    assert "client.email" in loose.variables


def test_raw_spanning_paragraphs(tmp_path):
    p = _doc(tmp_path, lambda d: [d.add_paragraph("{% raw %}{% if"), d.add_paragraph("x %}{% endraw %}{{ y }}")])
    rep = lint_path(p)
    assert rep.issues == [] and rep.variables == ["y"]


def test_ws_control_markers(tmp_path):
    p = _doc(tmp_path, lambda d: d.add_paragraph("{%- if a -%}x{%- endif -%}{{ b - 1 }}"))
    rep = lint_path(p)
    assert rep.issues == [] and rep.variables == ["a", "b"]


def test_else_outside_block(tmp_path):
    assert rules(lint_path(_doc(tmp_path, lambda d: d.add_paragraph("{% else %}")))) == ["DTL102"]


def test_other_syntax_error(tmp_path):
    assert rules(lint_path(_doc(tmp_path, lambda d: d.add_paragraph("{{ a | }}")))) == ["DTL103"]


def test_nested_table_location(tmp_path):
    def b(d):
        outer = d.add_table(rows=2, cols=2)
        outer.cell(1, 1).add_table(rows=2, cols=1).cell(1, 0).text = "{% if x %}"

    (i,) = lint_path(_doc(tmp_path, b)).issues
    assert i.location.cells == (CellPos(1, 2, 2), CellPos(2, 2, 1))
    assert "T1 R2 C2 > T2 R2 C1" in i.location.human()


def test_header_scanned(tmp_path):
    def b(d):
        d.sections[0].header.paragraphs[0].text = "{% if x %}"

    (i,) = lint_path(_doc(tmp_path, b)).issues
    assert i.location.part == "header1"


def test_comments_not_scanned(tmp_path):
    """docxtpl does not render comments.xml, so reviewer notes must not be linted."""

    def b(d):
        para = d.add_paragraph("x")
        d.add_comment(para.runs, text="{% if y %}")

    try:
        p = _doc(tmp_path, b)
    except AttributeError:
        pytest.skip("python-docx without comment support")
    assert lint_path(p).issues == []


def test_file_level_disable(tmp_path):
    p = _doc(tmp_path, lambda d: d.add_paragraph("{# docxtpl-lint: disable=DTL102 #}{% endif %}"))
    assert lint_path(p).issues == []


# --- variable analysis ---------------------------------------------------------


def test_scopes(tmp_path):
    def b(d):
        d.add_paragraph("{% with y = a.b %}{{ y.c }}{% endwith %}")
        d.add_paragraph("{% for k, v in m.items() %}{{ k }}{{ v }}{{ loop.index }}{% endfor %}")
        d.add_paragraph("{% for i in range(3) %}{{ i }}{% endfor %}")
        d.add_paragraph("{% macro f(z) %}{{ z }}{{ g }}{% endmacro %}{{ f(1) }}")
        d.add_paragraph("{% set s %}{{ h }}{% endset %}{{ s }}")

    rep = lint_path(_doc(tmp_path, b))
    assert rep.issues == []
    assert rep.variables == ["a.b", "g", "h", "m"]


def test_vars_listed_even_when_template_broken(fx):
    assert "client.name" in lint_path(fx / "issue378_unclosed.docx").variables


# --- data contract -------------------------------------------------------------


def test_data_contract(fx):
    data = json.loads((fx / "clean.data.json").read_text())
    rep = lint_path(fx / "clean.docx", data)
    assert [(i.rule, i.params.get("path")) for i in rep.issues] == [("DTL301", "items[*].price")]
    rep = lint_path(fx / "clean.docx", data, report_unused_data=True)
    assert ("DTL302", "unused_key") in [(i.rule, i.params.get("key")) for i in rep.issues]


def test_guards_not_reported(tmp_path):
    def b(d):
        d.add_paragraph("{% if note is defined %}{{ note.text }}{% endif %}")
        d.add_paragraph("{{ title | default('Untitled') }}")
        d.add_paragraph("{{ missing }}")

    rep = lint_path(_doc(tmp_path, b), {})
    assert [i.params["path"] for i in rep.issues] == ["missing"]


def test_empty_list_cannot_prove_absence(fx):
    assert lint_path(fx / "clean.docx", {"customer": {"name": 1}, "items": [], "company": 1}).issues == []


def test_multiple_contexts(tmp_path):
    p = _doc(tmp_path, lambda d: d.add_paragraph("{{ a }}{{ b }}{{ c }}"))
    rep = lint_path(p, contexts=[{"a": 1}, {"b": 1}])
    assert [i.params["path"] for i in rep.issues] == ["c"]


def test_non_iterable(tmp_path):
    p = _doc(tmp_path, lambda d: d.add_paragraph("{% for i in n %}{% endfor %}{% for c in s %}{% endfor %}"))
    rep = lint_path(p, {"n": 3, "s": "abc"})
    assert [(i.rule, i.params["path"], i.params["type"]) for i in rep.issues] == [("DTL303", "n", "number")]


# --- CLI -------------------------------------------------------------------------


def test_cli_exit_codes(fx):
    assert main([str(fx / "clean.docx")]) == 0
    assert main([str(fx / "issue644_percent.docx")]) == 1
    assert main([]) == 2
    assert main([str(fx / "nope.docx")]) == 2
    assert main([str(fx / "clean.docx"), "--data", str(fx / "missing.json")]) == 2


def test_cli_warnings_policy(fx, tmp_path):
    p = tmp_path / "w.docx"
    build_case(next(c for c in CASES if c.name == "cross_para_warn"), p)
    assert main([str(p)]) == 0
    assert main([str(p), "--fail-on", "warning"]) == 1
    assert main([str(p), "--warnings-as-errors"]) == 2  # removed in 2.0


def test_cli_select_ignore(fx):
    assert main([str(fx / "issue603_table_loop.docx"), "--ignore", "DTL2"]) == 0
    assert main([str(fx / "issue603_table_loop.docx"), "--select", "DTL1"]) == 0


def test_cli_json(fx, capsys):
    main([str(fx / "issue378_unclosed.docx"), "--format", "json"])
    out = json.loads(capsys.readouterr().out)
    assert out["schema_version"] == 1
    issue = next(i for i in out["files"][0]["issues"] if i["rule"] == "DTL101")
    assert issue["location"]["cells"] == [{"table": 2, "row": 3, "col": 1}]
    assert issue["message"].startswith("`{% if %}` is never closed")


def test_cli_lang_zh(fx, capsys):
    main([str(fx / "issue378_unclosed.docx"), "--lang", "zh", "--no-color"])
    out = capsys.readouterr().out
    assert "没有闭合" in out and "表2 行3 列1" in out and "is never closed" not in out


def test_cli_github(fx, capsys):
    assert main([str(fx / "issue644_percent.docx"), "--format", "github"]) == 1
    lines = capsys.readouterr().out.strip().splitlines()
    assert len(lines) == 2 and lines[0].startswith("::error file=")
    assert "title=DTL201 body ¶1" in lines[0]


def test_cli_sarif(fx, capsys):
    assert main([str(fx / "issue644_percent.docx"), "--format", "sarif"]) == 1
    out = json.loads(capsys.readouterr().out)
    res = out["runs"][0]["results"]
    assert out["version"] == "2.1.0" and [r["ruleId"] for r in res] == ["DTL201", "DTL201"]
    assert res[1]["locations"][0]["physicalLocation"]["region"]["startLine"] == 4


def test_cli_multi_data(fx, tmp_path):
    a, b = tmp_path / "a.json", tmp_path / "b.json"
    a.write_text(json.dumps({"customer": {"name": 1}, "items": [{"name": 1, "price": 1}]}))
    b.write_text(json.dumps({"company": 1}))
    assert main([str(fx / "clean.docx"), "--data", str(a)]) == 1
    assert main([str(fx / "clean.docx"), "--data", str(a), "--data", str(b)]) == 0


def test_cli_list_rules(capsys):
    assert main(["--list-rules"]) == 0
    assert capsys.readouterr().out.count("DTL") == len(RULES)


def test_cli_corrupt_file(tmp_path):
    bad = tmp_path / "bad.docx"
    bad.write_text("not a zip")
    assert main([str(bad)]) == 2


def test_never_crashes_on_examples(fx):
    for f in fx.glob("*.docx"):
        lint_path(f, {}, report_unused_data=True)


def test_integer_index(tmp_path):
    p = _doc(tmp_path, lambda d: d.add_paragraph("{{ items[0].name }}{{ items[0].qty }}"))
    rep = lint_path(p, {"items": [{"name": "a"}]}, report_unused_data=True)
    assert rep.variables == ["items[].name", "items[].qty"]
    assert [(i.rule, i.params.get("path")) for i in rep.issues] == [("DTL301", "items[*].qty")]
