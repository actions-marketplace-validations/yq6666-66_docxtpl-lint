"""v0.2 features: config file, --context, help URLs, benchmark script."""

import json
import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document

from docxtpl_lint.cli import main
from docxtpl_lint.config import ConfigError, find_config, load_config

ROOT = Path(__file__).resolve().parent.parent


def _doc(path, *texts):
    d = Document()
    for t in texts:
        d.add_paragraph(t)
    d.save(path)
    return path


def test_config_applies_and_cli_overrides(tmp_path, monkeypatch, capsys):
    (tmp_path / "pyproject.toml").write_text(
        '[tool.docxtpl-lint]\nignore = ["DTL102"]\ndata = ["ctx.json"]\nlang = "zh"\n', encoding="utf-8"
    )
    (tmp_path / "ctx.json").write_text(json.dumps({"a": 1}))
    sub = tmp_path / "templates"
    sub.mkdir()
    doc = _doc(sub / "t.docx", "{% endif %}{{ a }}{{ b }}")
    monkeypatch.chdir(sub)  # found by walking up
    assert main([str(doc), "--no-color"]) == 1
    out = capsys.readouterr().out
    assert "DTL102" not in out and "DTL301" in out and "数据文件里没有 `b`" in out
    assert main([str(doc), "--no-color", "--lang", "en"]) == 1
    assert "not present in the data file" in capsys.readouterr().out
    assert main([str(doc), "--no-config", "--no-color"]) == 1
    out = capsys.readouterr().out
    assert "DTL102" in out and "DTL301" not in out


def test_config_unknown_key_is_error(tmp_path, monkeypatch, capsys):
    (tmp_path / "pyproject.toml").write_text("[tool.docxtpl-lint]\nignor = ['DTL1']\n")
    doc = _doc(tmp_path / "t.docx", "x")
    monkeypatch.chdir(tmp_path)
    assert main([str(doc)]) == 2
    assert "unknown key" in capsys.readouterr().err
    with pytest.raises(ConfigError):
        load_config(tmp_path / "pyproject.toml")


def test_config_wrong_type(tmp_path):
    (tmp_path / "pyproject.toml").write_text("[tool.docxtpl-lint]\nfail-on = true\n")
    with pytest.raises(ConfigError):
        load_config(tmp_path / "pyproject.toml")


def test_find_config_ignores_pyproject_without_table(tmp_path):
    (tmp_path / "pyproject.toml").write_text("[project]\nname='x'\n")
    assert find_config(tmp_path) is None or find_config(tmp_path).parent != tmp_path


def test_context_lines(tmp_path, capsys):
    doc = _doc(tmp_path / "t.docx", "before", "{% if x %}", "after", "far")
    main([str(doc), "--no-color", "--context", "1", "--no-config"])
    out = capsys.readouterr().out
    assert "> ¶2" in out and "¶1" in out and "| before" in out and "| after" in out and "far" not in out


def test_context_default_off_when_not_a_tty(tmp_path, capsys):
    doc = _doc(tmp_path / "t.docx", "before", "{% if x %}")
    main([str(doc), "--no-color", "--no-config"])
    assert "| before" not in capsys.readouterr().out


def test_help_urls(tmp_path, capsys):
    doc = _doc(tmp_path / "t.docx", "{% if x %}")
    main([str(doc), "--format", "json", "--no-config"])
    issue = json.loads(capsys.readouterr().out)["files"][0]["issues"][0]
    assert issue["help_url"].endswith("docs/rules.md#dtl101")
    main([str(doc), "--format", "sarif", "--no-config"])
    rules = json.loads(capsys.readouterr().out)["runs"][0]["tool"]["driver"]["rules"]
    assert all(r["helpUri"].endswith("#" + r["id"].lower()) for r in rules)


def test_rules_doc_has_every_rule():
    from docxtpl_lint.messages import RULES

    doc = (ROOT / "docs" / "rules.md").read_text(encoding="utf-8")
    for rid in RULES:
        assert f'<a id="{rid.lower()}"></a>' in doc


def test_benchmark_script(tmp_path):
    """Mini benchmark: a clean template + the deliberately broken one."""
    from rule_cases import CASES, build

    tpl = tmp_path / "templates"
    tpl.mkdir()
    build(next(c for c in CASES if c.name == "tr_good"), tpl / "ok_tpl.docx")
    _doc(tpl / "template_error_tpl.docx", "{% endif %}")
    r = subprocess.run(
        [sys.executable, str(ROOT / "scripts" / "benchmark.py"), str(tpl)], capture_output=True, text=True
    )
    assert r.returncode == 0, r.stdout + r.stderr
    assert "false positives: 0" in r.stdout


def test_benchmark_flags_stale_allowlist(tmp_path):
    import importlib.util
    from pathlib import Path

    root = Path(__file__).resolve().parents[1]
    spec = importlib.util.spec_from_file_location("bench", root / "scripts/benchmark.py")
    bench = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(bench)
    res = bench.run(tmp_path, {"known_false_positives": {"gone.docx": ["DTL103"]}})
    assert res["stale_allowlist"] == []  # file absent -> nothing to compare
    import docx

    d = docx.Document()
    d.add_paragraph("Hello {{ name }}")
    d.save(tmp_path / "clean.docx")
    res = bench.run(tmp_path, {"known_false_positives": {"clean.docx": ["DTL103"]}})
    assert res["stale_allowlist"] == ["clean.docx:DTL103"]
    assert res["false_positives"] == []
